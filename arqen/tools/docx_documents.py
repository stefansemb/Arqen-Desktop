from pathlib import Path
from typing import Any

from arqen.tools.base import Tool
from arqen.config import paths


class ReadDocxTool(Tool):
    name = "read_docx"
    description = "Extracts text from a DOCX file inside the Arqen workspace."
    requires_confirmation = False
    arguments_schema = {"path": str}

    def run(self, arguments: dict[str, Any]) -> str:
        root = paths.workspace_root()
        candidate = (root / arguments["path"]).resolve()
        if root not in candidate.parents or candidate.suffix.lower() != ".docx":
            raise PermissionError("DOCX path must be inside the Arqen workspace")
        if not candidate.is_file():
            raise FileNotFoundError(f"File not found: {arguments['path']}")
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX reading requires the python-docx package") from exc

        document = Document(str(candidate))
        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts) or "No readable text found in the DOCX file."

