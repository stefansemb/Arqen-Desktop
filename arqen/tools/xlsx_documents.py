from pathlib import Path
from typing import Any
import zipfile
import xml.etree.ElementTree as ET

from arqen.tools.base import Tool


class ReadXlsxTool(Tool):
    name = "read_xlsx"
    description = "Reads sheet names and cell contents from an XLSX file inside the workspace."
    requires_confirmation = False
    arguments_schema = {"path": str}

    def run(self, arguments: dict[str, Any]) -> str:
        root = Path.cwd().resolve()
        candidate = (root / arguments["path"]).resolve()
        if root not in candidate.parents or candidate.suffix.lower() != ".xlsx":
            raise PermissionError("XLSX path must be inside the Arqen workspace")
        if not candidate.is_file():
            raise FileNotFoundError(f"File not found: {arguments['path']}")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX reading requires the openpyxl package") from exc

        try:
            workbook = load_workbook(candidate, read_only=True, data_only=True)
        except TypeError:
            return self._read_raw_xlsx(candidate)
        sections = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(max_row=100, max_col=30, values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    rows.append(" | ".join(values).rstrip(" |"))
            sections.append(f"Sheet: {sheet.title} :: " + (" || ".join(rows) or "(empty)"))
        workbook.close()
        return "\n\n".join(sections) or "Workbook is empty."

    @staticmethod
    def _read_raw_xlsx(path: Path) -> str:
        """Fallback reader that ignores style metadata from incompatible files."""
        ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with zipfile.ZipFile(path) as archive:
            shared = []
            if "xl/sharedStrings.xml" in archive.namelist():
                root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
                for item in root.findall("main:si", ns):
                    shared.append("".join(node.text or "" for node in item.iter() if node.tag.endswith("}t") or node.tag == "t"))
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            rel_map = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels}
            sections = []
            for sheet in workbook.findall("main:sheets/main:sheet", ns):
                name = sheet.attrib.get("name", "Sheet")
                rel_id = next(value for key, value in sheet.attrib.items() if key.endswith("}id") or key == "id")
                target = rel_map[rel_id]
                if not target.startswith("/"):
                    target = "xl/" + target.lstrip("/")
                root = ET.fromstring(archive.read(target))
                rows = []
                for row in root.findall(".//main:sheetData/main:row", ns):
                    values = []
                    for cell in row.findall("main:c", ns):
                        value = cell.find("main:v", ns)
                        text = value.text if value is not None else ""
                        if cell.attrib.get("t") == "s" and text.isdigit():
                            text = shared[int(text)]
                        values.append(text or "")
                    rows.append(" | ".join(values).rstrip(" |"))
                sections.append(f"Sheet: {name} :: " + (" || ".join(rows) or "(empty)"))
        return "\n\n".join(sections) or "Workbook is empty."
